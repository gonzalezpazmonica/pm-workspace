#!/usr/bin/env python3
"""SE-413/SE-415 — worker de extracción de Savia Files.

Uso:
  extract.py <pdf|docx|pptx|xlsx> <ruta>     un fichero; un objeto JSON en stdout
  extract.py --batch                        lote: una línea JSON {"type","path"} por
                                            fichero en stdin; una línea JSON de
                                            resultado por fichero en stdout, en orden
Resultado:
  {"method": str, "units": [{"locator": {...}, "kind": str, "text": str, "formula"?: str}],
   "skipped": [{"reason": str, "count": int}]}   o   {"error": str}

El contenido es dato no confiable: no se ejecutan macros, no se recalculan
fórmulas y no se hace OCR. Sin red: los modelos de Docling deben estar en caché.
En lote, el conversor de Docling se carga una sola vez (SE-415 E2) y cada fichero
tiene su propio límite de tiempo (SAVIA_FILES_ITEM_TIMEOUT_S).
"""
import io
import json
import logging
import os
import signal
import sys
import warnings
from collections import Counter

MAX_UNITS = int(os.environ.get("SAVIA_FILES_MAX_UNITS", "50000"))
MAX_UNIT_CHARS = 20000
ITEM_TIMEOUT_S = int(os.environ.get("SAVIA_FILES_ITEM_TIMEOUT_S", "300"))
TYPES = ("pdf", "docx", "pptx", "xlsx")

warnings.filterwarnings("ignore")
logging.disable(logging.WARNING)


class Out:
    def __init__(self, method):
        self.method = method
        self.units = []
        self.skipped = Counter()

    def add(self, locator, kind, text, formula=None):
        text = (text or "").strip()
        if not text:
            return
        if len(self.units) >= MAX_UNITS:
            self.skipped["max-units"] += 1
            return
        if len(text) > MAX_UNIT_CHARS:
            text = text[:MAX_UNIT_CHARS]
            self.skipped["truncated"] += 1
        unit = {"locator": locator, "kind": kind, "text": text}
        if formula:
            unit["formula"] = formula
        self.units.append(unit)

    def dump(self):
        return {
            "method": self.method,
            "units": self.units,
            "skipped": [{"reason": r, "count": c} for r, c in sorted(self.skipped.items())],
        }


def docling_version():
    try:
        from importlib.metadata import version
        return version("docling")
    except Exception:
        return "unknown"


_converter = None


def converter():
    """Conversor de Docling sin OCR; se crea una vez por proceso."""
    global _converter
    if _converter is None:
        from docling.datamodel.base_models import InputFormat
        from docling.datamodel.pipeline_options import PdfPipelineOptions
        from docling.document_converter import DocumentConverter, PdfFormatOption

        opts = PdfPipelineOptions(do_ocr=False)
        _converter = DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=opts)})
    return _converter


def pptx_notes(path, out):
    """SE-415 Q3: Docling no extrae las notas del presentador; python-pptx sí."""
    try:
        import pptx
    except ImportError:
        out.skipped["notes-unavailable"] += 1
        return
    with open(path, "rb") as fh:
        prs = pptx.Presentation(io.BytesIO(fh.read()))
    for number, slide in enumerate(prs.slides, start=1):
        if not slide.has_notes_slide:
            continue
        frame = slide.notes_slide.notes_text_frame
        text = frame.text if frame is not None else ""
        text = text.strip()
        if text:
            shown = text if text.lower().startswith("notas") else f"Notas: {text}"
            out.add({"type": "slide", "slide": number}, "notes", shown)


def docling_extract(kind, path):
    doc = converter().convert(path).document
    out = Out(f"docling-{docling_version()}")
    pages_with_text = set()
    for index, (item, _level) in enumerate(doc.iterate_items()):
        label = str(getattr(item, "label", "") or type(item).__name__).split(".")[-1].lower()
        text = getattr(item, "text", None)
        if text is None and label == "table":
            text = item.export_to_markdown(doc)
        if text is None:
            out.skipped["image-no-ocr" if label == "picture" else f"no-text-{label}"] += 1
            continue
        prov = getattr(item, "prov", None) or []
        page = prov[0].page_no if prov else None
        if kind == "pdf":
            locator = {"type": "page", "page": page or 1}
        elif kind == "pptx":
            locator = {"type": "slide", "slide": page or 1}
        else:
            locator = {"type": "element", "index": index}
        before = len(out.units)
        out.add(locator, label, text)
        if kind == "pdf" and len(out.units) > before:
            pages_with_text.add(page or 1)
    if kind == "pdf":
        # SE-415 Q1: páginas sin ninguna unidad de texto (escaneadas o solo imagen), declaradas.
        pages = set(getattr(doc, "pages", {}) or {})
        missing = len(pages - pages_with_text)
        if missing:
            out.skipped["page-without-text"] += missing
    if kind == "pptx":
        pptx_notes(path, out)
    return out


def header_row(rows):
    """Índice de la cabecera: la primera fila no vacía si tiene 2+ celdas, todas texto (SE-415 Q4)."""
    for r_index, row in enumerate(rows):
        values = [c for c in row if c is not None and str(c).strip() != ""]
        if not values:
            continue
        if len(values) >= 2 and all(isinstance(c, str) for c in values):
            return r_index
        return None
    return None


def xlsx_extract(path):
    import openpyxl
    from openpyxl.utils import get_column_letter

    # Los blobs se nombran por SHA-256, sin extensión: openpyxl valida la extensión
    # de una ruta, así que se le pasa un stream binario.
    with open(path, "rb") as fh:
        data = fh.read()
    formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True, keep_vba=False)
    values = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True, keep_vba=False)
    out = Out(f"openpyxl-{openpyxl.__version__}")
    for ws in formulas.worksheets:
        f_rows = list(ws.iter_rows(values_only=True))
        v_rows = list(values[ws.title].iter_rows(values_only=True))
        h = header_row(v_rows)
        headers = {}
        if h is not None:
            for c, v in enumerate(v_rows[h]):
                if isinstance(v, str) and v.strip():
                    headers[c] = v.strip()
        for r, (f_row, v_row) in enumerate(zip(f_rows, v_rows)):
            in_body = h is not None and r > h
            label = next((str(v).strip() for v in v_row if isinstance(v, str) and v.strip()), None) if in_body else None
            for c, raw in enumerate(f_row):
                if raw is None:
                    continue
                coord = f"{get_column_letter(c + 1)}{r + 1}"
                formula = raw if isinstance(raw, str) and raw.startswith("=") else None
                value = (v_row[c] if c < len(v_row) else None) if formula else raw
                if formula and value is None:
                    out.skipped["formula-without-cached-value"] += 1
                shown = "" if value is None else str(value)
                context = []
                if in_body:
                    if c in headers:
                        context.append(headers[c])
                    if label and label != shown.strip():
                        context.append(label)
                prefix = f"{ws.title}!{coord}" + "".join(f" · {x}" for x in context)
                text = f"{prefix}: {shown}" + (f" ({formula})" if formula else "")
                out.add({"type": "cell", "sheet": ws.title, "cell": coord}, "cell", text, formula)
    formulas.close()
    values.close()
    return out


class ItemTimeout(Exception):
    pass


def _on_alarm(_signum, _frame):
    raise ItemTimeout()


def extract(kind, path):
    try:
        out = xlsx_extract(path) if kind == "xlsx" else docling_extract(kind, path)
        return out.dump()
    except ItemTimeout:
        return {"error": f"timeout del worker ({ITEM_TIMEOUT_S} s por fichero)"}
    except Exception as e:  # el llamador marca FAILED con este mensaje
        return {"error": f"{type(e).__name__}: {str(e)[:500]}"}


def batch():
    signal.signal(signal.SIGALRM, _on_alarm)
    for line in sys.stdin:
        line = line.strip()
        if not line:
            continue
        try:
            item = json.loads(line)
            kind, path = item["type"], item["path"]
            if kind not in TYPES:
                raise ValueError(f"tipo no soportado: {kind}")
        except Exception as e:
            result = {"error": f"entrada no válida: {str(e)[:200]}"}
        else:
            signal.alarm(ITEM_TIMEOUT_S)
            try:
                result = extract(kind, path)
            finally:
                signal.alarm(0)
        sys.stdout.write(json.dumps(result, ensure_ascii=False) + "\n")
        sys.stdout.flush()
    return 0


def main(argv):
    if len(argv) == 2 and argv[1] == "--batch":
        return batch()
    if len(argv) != 3 or argv[1] not in TYPES:
        print(json.dumps({"error": "uso: extract.py <pdf|docx|pptx|xlsx> <ruta> | --batch"}))
        return 2
    result = extract(argv[1], argv[2])
    sys.stdout.write(json.dumps(result, ensure_ascii=False))
    return 1 if "error" in result else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
